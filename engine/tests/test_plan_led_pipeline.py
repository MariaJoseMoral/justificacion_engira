"""Focused regression tests for the plan-led justification foundation."""

from __future__ import annotations

import shutil
import sys
import unittest
import csv
from pathlib import Path
from uuid import uuid4

from docx import Document

ENGINE_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ENGINE_ROOT))

from src.evidencias import vincular_evidencias
from src.economica import normalizar_gasto_label
from src.memoria_actividades import GeneradorMemoriaActividades
from src.plan import CanonicalPlan, parse_plan_documents
from src.relacion_gastos import (
    InvoicePage,
    extraer_factura_desde_paginas,
    generar_relacion_gastos,
)


class PlanLedPipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.fixture_dir = Path(__file__).parent / f".fixture-{uuid4().hex}"
        self.fixture_dir.mkdir()
        self.chronogram = self.fixture_dir / "Cronograma aprobado.docx"
        self._write_chronogram(self.chronogram)
        self.plan = parse_plan_documents([self.chronogram], self.fixture_dir)

    def tearDown(self) -> None:
        shutil.rmtree(self.fixture_dir)

    def test_parses_stable_actions_and_phase_budget(self) -> None:
        repeated = parse_plan_documents([self.chronogram], self.fixture_dir)

        self.assertEqual(2, len(self.plan.actions))
        self.assertEqual("COMUNICACIÓN", self.plan.actions[0].phase)
        self.assertEqual("2025-07-01", self.plan.actions[0].period_start)
        self.assertEqual(2282.90, self.plan.phase_budgets["COMUNICACIÓN"])
        self.assertEqual(self.plan.actions[0].id, repeated.actions[0].id)

    def test_classifies_evidence_as_reviewable_action_link(self) -> None:
        links = vincular_evidencias(
            self.plan,
            [{
                "area": "comunicacion",
                "nombre": "campana_multicanal.pdf",
                "ruta_relativa": "04_COMUNICACION/campana_multicanal.pdf",
                "texto_preview": "Campaña multicanal para usuarios",
            }],
            ["comunicacion"],
        )

        self.assertEqual(1, len(links))
        self.assertEqual(self.plan.actions[0].id, links[0].action_id)
        self.assertEqual("automatico", links[0].review_state)
        self.assertGreaterEqual(links[0].confidence, 0.75)

    def test_generates_activity_docx_from_actions_and_evidence_paths(self) -> None:
        template = self.fixture_dir / "plantilla.docx"
        output = self.fixture_dir / "salida.docx"
        document = Document()
        table = document.add_table(rows=3, cols=4)
        table.rows[0].cells[0].text = "Fase"
        document.save(template)
        links = vincular_evidencias(
            self.plan,
            [{
                "area": "comunicacion",
                "nombre": "campana_multicanal.pdf",
                "ruta_relativa": "04_COMUNICACION/campana_multicanal.pdf",
                "texto_preview": "Campaña multicanal",
            }],
            ["comunicacion"],
        )

        GeneradorMemoriaActividades({"proyecto": {}}, self.plan, links).generar(template, output)

        result = Document(output).tables[0]
        self.assertIn("COMUNICACIÓN", result.rows[1].cells[0].text)
        self.assertIn(self.plan.actions[0].id, result.rows[1].cells[2].text)
        self.assertIn("04_COMUNICACION/campana_multicanal.pdf", result.rows[1].cells[2].text)
        self.assertNotIn("No se ha extraído texto", result.rows[1].cells[2].text)

    @staticmethod
    def _write_chronogram(path: Path) -> None:
        document = Document()
        table = document.add_table(rows=4, cols=4)
        headers = (
            "DENOMINACIÓN DE LA FASE",
            "FECHAS PREVISTAS",
            "PRINCIPALES ACTUACIONES A REALIZAR",
            "CUANTÍA DEL GASTO",
        )
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
        rows = (
            ("COMUNICACIÓN", "Julio 2025 – Junio 2026", "Campaña multicanal", "2.282,90 €"),
            ("COMUNICACIÓN", "Julio 2025 – Junio 2026", "Newsletter usuarios", "2.282,90 €"),
            ("SUMA TOTAL", "", "", "2.282,90 €"),
        )
        for row, values in zip(table.rows[1:], rows):
            for cell, value in zip(row.cells, values):
                cell.text = value
        document.save(path)


class InvoiceExtractionTests(unittest.TestCase):
    def test_normalizes_gasto_variants_to_shared_label(self) -> None:
        self.assertEqual("externos", normalizar_gasto_label("proveedores externos"))
        self.assertEqual("externos", normalizar_gasto_label("Externo"))
        self.assertEqual("autónoma", normalizar_gasto_label("Trabajos realizados por el autónomo"))
        self.assertEqual("nómina", normalizar_gasto_label("Nóminas y personal"))

    def test_extracts_labelled_parties_dates_and_table_amounts(self) -> None:
        extraction = extraer_factura_desde_paginas((
            InvoicePage(
                1,
                "\n".join((
                    "FACTURA Nº FV-2025-42",
                    "Proveedor: Proveedora Teatro, S.L.",
                    "Cliente: enGira, S.L.",
                    "Fecha de emisión: 14/10/2025",
                    "Fecha de vencimiento: 30/10/2025",
                    "Retención IRPF: 15,00 EUR",
                )),
                (
                    (
                        ("Descripción", "Base imponible", "IVA", "Total"),
                        ("Servicio de mediación", "100,00 €", "21,00 €", "106,00 €"),
                    ),
                ),
            ),
        ))

        self.assertEqual("Proveedora Teatro, S.L.", extraction.value("emisor"))
        self.assertEqual("enGira, S.L.", extraction.value("receptor"))
        self.assertEqual("FV-2025-42", extraction.value("numero_factura"))
        self.assertEqual("2025-10-14", extraction.value("fecha_emision"))
        self.assertEqual("2025-10-30", extraction.value("fecha_vencimiento"))
        self.assertEqual("100.00", extraction.value("base_imponible"))
        self.assertEqual("21.00", extraction.value("iva"))
        self.assertEqual("15.00", extraction.value("retencion"))
        self.assertEqual("106.00", extraction.value("total"))
        self.assertEqual("Servicio de mediación", extraction.value("concepto"))

        base = next(item for item in extraction.provenance if item.field == "base_imponible")
        self.assertEqual(1, base.page)
        self.assertEqual("tabla", base.source_type)
        self.assertFalse(base.ocr_needed)

    def test_marks_native_text_absence_as_ocr_review_without_inventing_values(self) -> None:
        extraction = extraer_factura_desde_paginas((InvoicePage(1),))

        self.assertTrue(extraction.ocr_needed)
        self.assertEqual("", extraction.value("emisor"))
        self.assertEqual("", extraction.value("total"))
        self.assertTrue(all(item.ocr_needed for item in extraction.provenance))
        self.assertTrue(all(item.review_status == "requiere_ocr" for item in extraction.provenance))

    def test_distinguishes_data_blocks_for_supplier_and_customer(self) -> None:
        extraction = extraer_factura_desde_paginas((
            InvoicePage(
                1,
                "\n".join((
                    "\uf23c\uf23c Datos Empresa \uf23c\uf23c Datos Factura",
                    "LABORAL RGPD, S.L.U. - B25857798",
                    "CORREGIDOR ESCOFET, 50",
                    "25005 - LLEIDA Fecha 14/10/2025",
                    "Datos Cliente 15554 - 26",
                    "Empresa MARIA JOSE MORAL MORGADO NIF 46943079X",
                )),
            ),
        ))

        self.assertEqual("LABORAL RGPD, S.L.U. - B25857798", extraction.value("emisor"))
        self.assertEqual("Empresa MARIA JOSE MORAL MORGADO NIF 46943079X", extraction.value("receptor"))
        self.assertEqual("2025-10-14", extraction.value("fecha_emision"))

    def test_prefers_aggregate_row_in_amount_table(self) -> None:
        extraction = extraer_factura_desde_paginas((
            InvoicePage(
                1,
                tables=((
                    ("Descripción", "Base imponible", "IVA", "Total"),
                    ("Servicio", "100,00 €", "21,00 €", "121,00 €"),
                    ("TOTAL", "120,00 €", "25,20 €", "145,20 €"),
                ),),
            ),
        ))

        self.assertEqual("120.00", extraction.value("base_imponible"))
        self.assertEqual("25.20", extraction.value("iva"))
        self.assertEqual("145.20", extraction.value("total"))

    def test_reads_summary_when_pdf_table_is_only_native_text(self) -> None:
        extraction = extraer_factura_desde_paginas((
            InvoicePage(
                1,
                "\n".join((
                    "Artículo Descripción Unidades B.I. IVA Total",
                    "Servicio 1 100,00 € 21,00 € 121,00 €",
                    "TOTAL 120,00 € 25,20 € 145,20 €",
                )),
            ),
        ))

        self.assertEqual("120.00", extraction.value("base_imponible"))
        self.assertEqual("25.20", extraction.value("iva"))
        self.assertEqual("145.20", extraction.value("total"))

    def test_keeps_main_schema_and_writes_field_audit_csv(self) -> None:
        fixture_dir = Path(__file__).parent / f".invoice-fixture-{uuid4().hex}"
        fixture_dir.mkdir()
        try:
            invoice = fixture_dir / "factura.docx"
            document = Document()
            document.add_paragraph("Proveedor: Proveedora Teatro, S.L.")
            document.add_paragraph("Cliente: enGira, S.L.")
            document.add_paragraph("Fecha de emisión: 14/10/2025")
            table = document.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "Descripción"
            table.rows[0].cells[1].text = "Base imponible"
            table.rows[1].cells[0].text = "Servicio de mediación"
            table.rows[1].cells[1].text = "100,00 €"
            document.save(invoice)
            output = fixture_dir / "relacion.csv"
            fields = fixture_dir / "campos.csv"
            review = fixture_dir / "revision.csv"

            relation = generar_relacion_gastos(
                [{
                    "ruta_relativa": "08/CONTABILIDAD/GASTOS/FACTURAS/factura.docx",
                    "ruta_absoluta": str(invoice),
                    "nombre": invoice.name,
                    "tipo_documental": "factura",
                }],
                CanonicalPlan((), (), {}),
                str(output),
                str(fields),
                str(review),
            )

            self.assertEqual(
                [
                    "id", "nombre_archivo", "emisor", "partida", "concepto_de_gasto",
                    "fecha_emision", "fecha_pago", "cargo_subvencion", "cargo_otros",
                    "importe_sin_iva", "nomina", "factura",
                ],
                list(relation.columns),
            )
            self.assertEqual("", relation.loc[0, "fecha_pago"])
            self.assertEqual("", relation.loc[0, "cargo_subvencion"])
            with fields.open(encoding="utf-8", newline="") as file:
                records = list(csv.DictReader(file))
            self.assertTrue(any(record["campo"] == "emisor" and record["pagina"] == "1" for record in records))
            self.assertTrue(any(record["campo"] == "fecha_pago" and not record["valor"] for record in records))
        finally:
            shutil.rmtree(fixture_dir)


if __name__ == "__main__":
    unittest.main()
